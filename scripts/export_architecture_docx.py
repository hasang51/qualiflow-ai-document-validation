"""Export docs/QualiFlow_Mimari_Genel_Bakis.md content to Word (.docx)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "QualiFlow_Mimari_Genel_Bakis.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def add_code_block(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val
    doc.add_paragraph()


def build() -> None:
    doc = Document()

    title = doc.add_heading("QualiFlow — Yüksek Seviyeli Mimari (Genel Bakış)", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_para(
        doc,
        "Proje: Endüstriyel Kalite Belgesi (CoA / MTC) Doğrulama Sistemi | "
        "Mimari: Neuro-symbolic (Yapay zeka okur + Kurallar doğrular)",
    )
    doc.add_paragraph()

    add_heading(doc, "1. Projenin Amacı", 1)
    add_para(
        doc,
        "QualiFlow, fabrikalara gelen Analiz Sertifikası (CoA) ve Değirmen Test Sertifikası (MTC) "
        "PDF belgelerinden otomatik veri çıkarır ve malzeme spesifikasyonlarına karşı uyumluluk kontrolü yapar. "
        "Emin olamadığı durumlarda yanlış onay vermek yerine insan incelemesine yönlendirir.",
    )

    add_heading(doc, "2. Sistem Bileşenleri (3 Katman)", 1)
    add_table(
        doc,
        ["Katman", "Klasör", "Görev"],
        [
            ["Arayüz", "frontend/", "PDF yükleme, sonuç ve inceleme nedenlerini gösterme"],
            ["Sunucu", "app/", "Profil → yönlendirme → çıkarım → doğrulama"],
            ["Veri", "data/", "Veritabanı, PDF dosyaları, test ve eval sonuçları"],
        ],
    )
    add_para(doc, "Dış bağımlılıklar: Amazon Bedrock Gemma (belge okuma), Poppler (PDF→görüntü).")

    add_heading(doc, "3. Genel Mimari Şema", 1)
    add_code_block(
        doc,
        "Kullanıcı → frontend (React) → app (FastAPI) → data (SQLite + PDF)\n"
        "                              ↓\n"
        "                         Amazon Bedrock Gemma",
    )

    add_heading(doc, "4. Belge İşleme Akışı (7 Adım)", 1)
    steps = [
        "① Profilleme — belge kalitesi (dijital / tarama / bozuk)",
        "② Yönlendirme — tek işleme yolu seçimi",
        "③ Ön işleme — PDF sayfalarını görüntüye çevirme",
        "④ Çıkarım (Gemma) — üst bilgi + tablo satırları",
        "⑤ Normalizasyon — alan adları standart forma",
        "⑥ Doğrulama — grade/spec uyum kontrolü",
        "⑦ Güven + inceleme — otomatik onay veya NEEDS_REVIEW",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Number")

    add_heading(doc, "5. Neuro-Symbolic Ayrımı", 1)
    add_table(
        doc,
        ["Bölüm", "Teknoloji", "Görev"],
        [
            ["Neural", "Bedrock Gemma multimodal", "Belgeyi okur, alanları çıkarır"],
            ["Symbolic", "Python kuralları", "Grade/spec kontrolü, review kararı"],
        ],
    )

    add_heading(doc, "6. Kalite Sınıfı ve Yönlendirme", 1)
    add_table(
        doc,
        ["Kalite", "Anlam", "Yol"],
        [
            ["digital_clean", "Dijital PDF", "Doğrudan multimodal"],
            ["scan_clean", "Temiz tarama", "Raster + multimodal"],
            ["noisy_scan", "Bozuk tarama", "Ön işlemeli multimodal"],
            ["severe_scan", "Çok kötü tarama", "Ön işlemeli + dikkatli inceleme"],
        ],
    )

    add_heading(doc, "7. Klasör Yapısı", 1)
    add_code_block(
        doc,
        "app/       → Backend (routes, services, domain)\n"
        "frontend/  → React arayüzü\n"
        "data/      → DB, PDF, gold set, batch sonuçları\n"
        "scripts/   → Toplu eval araçları\n"
        "tests/     → pytest (~240 test)",
    )

    add_heading(doc, "8. Güvenlik (Özet)", 1)
    for item in [
        "API anahtarları .env dosyasında; Git'e gönderilmez.",
        "Kullanıcı şifreleri bcrypt hash ile veritabanında saklanır.",
        "Bedrock erişimi IAM varsayılan kimlik zinciri ile sunucuda kullanılır.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "9. Özet", 1)
    add_para(
        doc,
        "QualiFlow, kalite belgelerini hibrit pipeline ile işler; yapay zeka okuma yapar, "
        "deterministik kurallar doğrular ve belirsizlikte açık inceleme gerekçeleri üretir.",
        bold=True,
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
